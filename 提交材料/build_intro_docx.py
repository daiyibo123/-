# -*- coding: utf-8 -*-
"""生成《手搓神经网络》作品介绍文档 docx（python-docx 路线）。"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "手搓神经网络_作品介绍文档.docx"
FIG_DIR = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs", "交付素材", "截图"))

doc = Document()

# ---------- 基础样式 ----------
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")

h1 = doc.styles["Heading 1"]
h1.font.name = "Times New Roman"
h1.font.size = Pt(15)
h1.font.bold = True
h1.font.color.rgb = RGBColor(0, 0, 0)
h1.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")
h1.paragraph_format.space_before = Pt(14)
h1.paragraph_format.space_after = Pt(7)

sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)

def set_ea(run):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")

def body_para(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.33)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.size = Pt(12)
    set_ea(r)
    return p

def center_para(text, size, bold=False, before=0, after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    set_ea(r)
    return p

# ---------- 页眉页脚 ----------
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hr = hp.add_run("《手搓神经网络》作品介绍文档")
hr.font.size = Pt(9)
set_ea(hr)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld_b = OxmlElement("w:fldChar"); fld_b.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
fld_e = OxmlElement("w:fldChar"); fld_e.set(qn("w:fldCharType"), "end")
fr = fp.add_run(); fr.font.size = Pt(9); set_ea(fr)
fr._element.append(fld_b)
fr._element.append(instr)
fr._element.append(fld_e)

def add_figure(filename, caption, width_in=5.5):
    path = os.path.join(FIG_DIR, filename)
    doc.add_picture(path, width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(12)
    r = cp.add_run(caption)
    r.font.size = Pt(10.5)
    set_ea(r)

# ---------- 封面 ----------
center_para("手搓神经网络", 22, bold=True, before=120, after=6)
center_para("深度学习课程交互式学习系统", 14, after=40)
center_para("作品介绍文档", 16, bold=True, after=70)
for line in [
    "参赛赛事　2026 年（第十四届）全国大学生数字媒体科技作品及创意竞赛",
    "作品类型　Web 端 H5 交互学习应用，浏览器即开即玩，免安装免注册",
    "团队名称　【待填】",
    "团队成员　【待填，共五人】",
    "指导教师　【待填，一至两人】",
    "申报院校　长春理工大学",
    "完成时间　2026 年 10 月",
]:
    center_para(line, 12, after=4)
pb = doc.add_paragraph()
pb.add_run().add_break(WD_BREAK.PAGE)

# ---------- 目录（TOC 域 + 缓存条目） ----------
toc_title = doc.add_paragraph()
tr = toc_title.add_run("目录"); tr.font.size = Pt(15); tr.font.bold = True; set_ea(tr)

p_begin = doc.add_paragraph()
r1 = p_begin.add_run()
fc1 = OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"), "begin"); fc1.set(qn("w:dirty"), "true")
it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "1-1" \\h \\z \\u '
fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "separate")
r1._element.append(fc1); r1._element.append(it); r1._element.append(fc2)

TOC_ENTRIES = [
    ("一、作品概述", "2"),
    ("二、创意来源", "2"),
    ("三、设计原则", "3"),
    ("四、创意与技术突破", "4"),
    ("五、功能与内容体系", "5"),
    ("六、运行与操作方法", "6"),
    ("七、试用情况与迭代", "6"),
    ("八、应用场景与推广价值", "7"),
    ("九、声明", "7"),
]
for title, page in TOC_ENTRIES:
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r = p.add_run(title); r.font.size = Pt(12); set_ea(r)
    r2 = p.add_run("\t" + page); r2.font.size = Pt(12); set_ea(r2)

p_end = doc.add_paragraph()
re_ = p_end.add_run()
fc3 = OxmlElement("w:fldChar"); fc3.set(qn("w:fldCharType"), "end")
re_._element.append(fc3)
p_end.add_run().add_break(WD_BREAK.PAGE)

# ---------- 正文 ----------
SECTIONS = [
("一、作品概述", [
"《手搓神经网络》是我们团队为高校深度学习课程开发的一款网页闯关学习系统。学生在浏览器里打开即可使用，不需要安装软件，也不需要注册账号。它和常见的教学演示工具不同，学生不是观看神经网络的训练动画，而是亲手扮演网络里的各个部件。学决策边界时直接拖动直线端点，看误分类点实时闪烁。学梯度下降时用摇杆选下山方向，用拨杆换挡调学习率，步子迈大了可以亲眼看到损失值爆表。学反向传播时把梯度卡片投进对应权重的槽位，亲手把误差从输出端一站一站送回去。学卷积时拖着滑窗在图片上逐格滑动，看着自己算出的特征图一点点亮起来。学生的每个操作都对应课程公式里的一个具体步骤，做错时内置的 AI 助教突触不会直接报答案，而是用追问引导学生发现矛盾，完成一章后系统生成个人学情报告，把薄弱点对应到教材章节。",
"作品覆盖五章课程内容，每章四个互动案例，主线共 20 个案例，另有 27 张知识图鉴卡牌和一条先修基础线。系统已在本校组织试玩，28 名同学完成了主要关卡的体验，主线通关率在 60% 到 80% 之间。下文依次介绍这个作品的创意来源、设计原则、技术突破和内容体系。",
("img", "01_课程路径页.png", "图 1　课程路径页，五章内容以路径形式组织", 5.5),
]),
("二、创意来源", [
"做这个作品，起初是因为深度学习教学里有几个多年解决不了的问题。第一个是数学黑箱感。前向传播、反向传播、链式法则，学生在课堂上跟着推导似乎能看懂，合上书还是无法在具体数值上走通一遍。第二个是训练直觉缺失。学习率不能太大、动量可以加速收敛、正则化防止过拟合，这些结论学生大多是背下来的。没有亲手把损失函数炸掉过一次，就很难真正理解为什么。第三个是实验成本高。真跑深度学习实验要配环境、要算力，课上时间有限，很难让每个人都动手调一遍。",
"我们调研了现有工具，确认了这个空位。TensorFlow Playground 让学生在浏览器里调超参数看训练过程，但学生从不亲手计算任何一步。CNN Explainer 和 Transformer Explainer 做得非常精致，本质仍是可交互的阅读材料。3Blue1Brown 的视频讲得很清楚，可那是单向传递，没有作答，也没有个体反馈。于是我们想做一件还没有人做成的事，让学生亲手回传一次误差，亲手当一次卷积核，亲手连出注意力权重，做完之后还能拿到一份基于错误归因的学情数据。",
]),
("三、设计原则", [
"作品开发全程守住两条设计律，团队评审任何新功能都用这两条来裁决。第一条是动作必须等于公式的一步，学生的每个操作都要能对应课程公式或伪代码里的一个具体步骤，不允许包装成与知识点无关的小游戏。第二条是交互动词必须多样，除点击选项外，每一关都要以拖拽、旋转、连线、摇杆、拼装、标尺、拉杆这类肢体化交互为主要操作，任何一关退化成点选项就判不合格。这两条原则保证了作品是一件教学工具而不是游戏外壳，手感上的新鲜感也来自这里。",
"在这种约束下，每个案例长成它要教的概念的样子。教决策边界的关卡，画面就是一条可以拖动的直线和一片会闪烁的点阵。教梯度下降的关卡，画面就是一座由损失曲面构成的山，学生摇着摇杆下山。教反向传播的关卡，画面是一条流水线，误差被包成包裹逐站送回。教注意力的关卡，学生从单词之间拖出连线调整粗细，再观察 softmax 归一化前后线宽的变化。每章的教学流程分四段，先是一段精讲翻卡交代本章重点和难点预告，然后是按序解锁的四个案例，接着是知识路智能体以苏格拉底问答的方式逐节点点灯，最后是单元测评和一份个人分析报告。上章测评过关才能解锁下一章，报告由规则引擎生成，大纲要求逐条对照，分数由案例星级、测评首答和知识路表现三部分合成，并附误区暴露榜前三和学习行为分析。",
("img", "03_注意力手算台.png", "图 2　注意力手算台案例的交互画面", 5.5),
]),
("四、创意与技术突破", [
"这个作品最核心的创意，是把深度学习的训练过程从观看变成操作。已有的流行工具全部停留在可视化层面，学生看动画、看曲线、看注意力热图，始终是旁观者。本作品让学生亲手完成一次完整的反向传播投递，亲手设计一枚卷积核并看着特征图随拨数实时变化，亲手拖出注意力连线并观察 softmax 归一化前后的线宽重排。据我们调研，还没有同类学习工具做到这个程度，学生的手在这里代替了自动求导引擎，这是作品区别于一切已有工具的根本所在。",
"第二个突破是页面内的真实计算。多数教学演示使用预录制动画或手绘曲线，学生看到的数字无法追问来源。本作品中 GAN 案例的生成器和判别器在浏览器里真实交替训练，每一轮的损失和生成样本都是现场算出来的，模式崩溃现象也是真实训练里自发出现的。反向传播案例的梯度值由离线脚本从固定种子生成，并用有限差分做了双重校验，误差在万分之一以内，页面上的梯度检查按钮会在运行时重新计算有限差分，再与学生投递的答案逐项对比。训练曲线来自真实小实验的离线导出并随代码版本存档。个别属于教学构造的演示曲线在页面上作了明确标注，不伪装成真实数据，这种对数值可验证性的坚持在教学软件里并不多见。",
("img", "06_GAN真实训练.png", "图 3　GAN 案例在页面内进行真实的生成器判别器交替训练", 5.5),
"第三个突破是贯穿全系统的错误归因引擎。我们整理了 49 条常见误区并编号入库，每张干扰卡片和每种非法操作都预绑定一个误区标签，学生每次犯错都会被打上标签并写入动作流。同一份数据贯穿提示阶梯、AI 助教追问、结算页归因卡和学情报告四个环节，学生知道自己卡在哪一章哪一页，教师从班级看板看到全班普遍没理解的知识点排行。交互过程直接转成可归因的学情数据，把学习行为测量做进了游戏流程里，这是传统课件和演示工具都没有做到的。",
"第四个突破是 AI 助教突触与三级降级链的设计。突触是一只纯 SVG 手绘的宠物形象，常驻在各个界面，学生操作被判错时它会走到出错位置旁边主动开口，但每次只给一个提示，学生重复犯错才逐步升级，绝不直接报数值答案，这种苏格拉底式的追问方式比直接讲解更能暴露学生的真实困惑。AI 服务按三级降级设计，优先调用后端接口，失败时查本地缓存，再失败时使用内置提示阶梯，断网也能完整游玩，这在课堂和比赛演示环境中是关键保障。AI 生成内容在对话流中固定标注来源，匿名会话限流，提问记录留存以备合规审计。",
("img", "07_AI助手突触_640.png", "图 4　AI 助教突触在各界面常驻", 3.2),
"第五个突破在工程架构上。前端使用原生 ES Modules，配合 SVG 和 Canvas 绘制界面，不依赖任何框架和构建工具，整个 webapp 目录就是一个静态站点，放到任意静态服务器即可部署。交互全部基于 Pointer Events，鼠标和触摸统一处理，触控目标不小于 44 像素。进度数据存在浏览器自带的 IndexedDB 里，匿名会话不需要登录，关闭上报开关后数据只留在本地。内容与代码分离，关卡数值、知识点和提示语全部放在 JSON 数据里，换一套知识点数据就可以迁移到其他人工智能课程或岗位培训。后端是可选项，用 Node.js 20 加 Express 加 better-sqlite3 实现，是一个单文件小服务，提供 AI 代理、事件上报和班级看板三个接口，密钥只存服务端。",
]),
("五、功能与内容体系", [
"作品内容对齐长春理工大学深度学习课程的教学大纲，知识点引用来自授课教师提供的课程资料。主线共五章，每章四个案例。第二章深度学习基础，含划一刀、误差往回流、梯度下山、诊室先预测四个案例，针对决策边界与偏置、模块化反向传播、学习率与动量、数据集划分与过拟合。第三章卷积神经网络，含尺寸车间、核工坊、平移的阴谋、结构进化史，覆盖输出尺寸与参数量计算、卷积核设计、池化不变性和从 LeNet 到 ResNet 的结构演进。第四章序列模型，含展开时间轴、梯度的长途旅行、三扇门、一句一句生成，覆盖权重共享、梯度消失、LSTM 门控和序列生成。第五章 Transformer，含手算台、读心连线、因果之幕、王座之争，覆盖注意力手算、温度与多头、因果掩码和位置编码。第六章生成模型，含压缩工坊、潜空间、猫鼠大战、从噪声里捞图，覆盖自编码器、变分自编码器、GAN 和扩散模型。",
"主线之外，系统配有 27 张知识图鉴卡牌，其中 14 张章节卡随通关解锁，13 张先修基础卡默认开放，覆盖线性代数、微积分、概率和机器学习入门概念，方便基础薄弱的学生先补先修。错题归因方面，49 条误区标签同时驱动结算页的归因卡和个人报告里的误区暴露榜，图鉴卡牌背面附有公式和课件章节页码引用，复习时可以顺着报告直接翻回教材。",
("img", "02_精讲翻卡.png", "图 5　每章开场的精讲翻卡", 5.5),
("img", "04_知识路智能体.png", "图 6　知识路智能体的苏格拉底问答界面", 5.5),
("img", "05_个人分析报告.png", "图 7　锚定课程大纲的个人分析报告", 5.5),
]),
("六、运行与操作方法", [
"本地运行只需双击启动手搓神经网络批处理文件，浏览器会自动打开本地 8799 端口对应的页面。等价方式是装好 Node.js 18 以上版本后运行 serve 脚本并指定端口 8799，全程不需要安装任何依赖。部署时 webapp 目录本身就是静态站点，放到任意静态服务器即可，服务器端组件按需部署以启用 AI 和数据采集。我们同时准备离线压缩包作为演示日的备份。",
"评审老师如果试玩，推荐这样一条路径。开场剧情之后进入课程路径页，选第二章，先翻精讲卡，再依次挑战四个案例。误差往回流这一关最能体现作品特点，值得完整体验一次纠错过程，包括被判定错误后突触走到出错位置追问的完整演出。之后可以走一遍知识路问答和单元测评，最后查看系统生成的个人分析报告。完整流程约四十到六十分钟，单个案例五到八分钟，也适合碎片化体验。",
]),
("七、试用情况与迭代", [
"作品完成后我们组织了试玩。28 名同学体验了主要关卡，主线 20 个案例的整体通关率在 60% 到 80% 之间。试玩反馈直接驱动了五轮可用性迭代，包括视觉体系的三版演进、案例页面全部改成单屏内完成、SVG 字号按四档标准化、AI 面板布局修正等，每一处改动都有具体的试玩反馈作为依据。测试过程中还修复了六处教学判定缺陷，并逐一做了回归验证。目前全部案例页都能在单屏内完成操作，不会因滚动遗漏任何交互，这一屏约束在移动端和投屏演示时尤其重要。",
]),
("八、应用场景与推广价值", [
"这个作品首先服务选课学生的课后自学，把复习变成闯关。教师也可以把单个案例投屏当教具，比如用误差往回流这一关在课堂上演示一次完整的反向传播。部署后端后，教师能看到班级匿名聚合的易错点排行，据此调整讲课重点。由于内容和代码解耦，换一套知识点数据就可以迁移到其他人工智能课程或岗位培训，推广成本很低。作品全程匿名采集、本地优先的设计也让它可以进入对隐私要求较高的教学环境。",
]),
("九、声明", [
"本作品为参赛团队独立设计开发的原创成果。作品中涉及的第三方素材均已注明来源，未使用存在版权争议的内容。开发过程中使用了人工智能工具辅助代码编写，全部生成代码经团队成员阅读理解并通过测试验收。AI 助教的行为规则和提示模板固化在服务端并留有调用日志。人工智能工具使用情况的详细说明随本材料一并提交。",
]),
]

for title, paras in SECTIONS:
    doc.add_heading(title, level=1)
    for item in paras:
        if isinstance(item, tuple) and item[0] == "img":
            add_figure(item[1], item[2], item[3])
        else:
            body_para(item)

doc.save(OUT)
print("saved:", OUT)
